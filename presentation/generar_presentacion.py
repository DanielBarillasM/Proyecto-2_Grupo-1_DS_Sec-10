"""Genera la presentación PowerPoint del EDA de MAP a partir de activos locales."""

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "report" / "assets"
OUTPUT = Path(__file__).with_name("presentacion_MAP.pptx")

NAVY = RGBColor(7, 23, 46)
BLUE = RGBColor(13, 71, 161)
CYAN = RGBColor(0, 174, 239)
GOLD = RGBColor(245, 181, 46)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(28, 42, 61)
MUTED = RGBColor(91, 107, 126)
PALE = RGBColor(240, 246, 252)
GREEN = RGBColor(29, 150, 108)
RED = RGBColor(207, 70, 66)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    p.text = value
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def rich_text(slide, parts, x, y, w, h, size=20, color=INK, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for i, item in enumerate(parts):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(11)
        if bullet:
            p.text = "•  " + p.text
    return box


def base_slide(title_value, kicker=None, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, WHITE)
    rect(slide, 0, 0, 0.18, 7.5, CYAN)
    if kicker:
        text(slide, kicker.upper(), 0.62, 0.32, 6.5, 0.28, 10, BLUE, True)
    text(slide, title_value, 0.62, 0.66, 11.9, 0.7, 26, NAVY, True)
    rect(slide, 0.62, 1.39, 1.05, 0.05, GOLD)
    text(slide, "MAP · Análisis Exploratorio · UVG", 0.62, 7.13, 6.0, 0.2, 8, MUTED)
    if number is not None:
        text(slide, f"{number:02d}", 12.15, 7.08, 0.55, 0.25, 9, MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, title_value, body, accent=CYAN, body_size=15):
    rect(slide, x, y, w, h, PALE, True, RGBColor(218, 228, 239))
    rect(slide, x, y, 0.07, h, accent, True, accent)
    text(slide, title_value, x + 0.25, y + 0.18, w - 0.45, 0.35, 15, NAVY, True)
    text(slide, body, x + 0.25, y + 0.65, w - 0.45, h - 0.82, body_size, MUTED)


def metric(slide, x, y, w, value, label, accent=CYAN):
    rect(slide, x, y, w, 1.18, WHITE, True, RGBColor(215, 226, 238))
    text(slide, value, x + 0.15, y + 0.12, w - 0.3, 0.5, 24, accent, True, align=PP_ALIGN.CENTER)
    text(slide, label, x + 0.12, y + 0.68, w - 0.24, 0.28, 10, MUTED, True, align=PP_ALIGN.CENTER)


def picture(slide, name, x, y, w, h=None):
    path = ASSETS / name
    if h is None:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    else:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


# 1 · Portada
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 0, 0.22, 7.5, CYAN)
rect(s, 0.82, 0.6, 1.2, 0.06, GOLD)
picture(s, "logo_uvg.png", 11.18, 0.52, 0.95, 0.86)
text(s, "ANÁLISIS EXPLORATORIO · PROYECTO 2", 0.82, 0.87, 8.7, 0.35, 12, CYAN, True)
text(s, "Charting Student Math\nMisunderstandings", 0.82, 1.35, 10.8, 1.45, 34, WHITE, True)
text(s, "Procesamiento del lenguaje natural para identificar patrones de razonamiento matemático", 0.84, 3.02, 9.4, 0.75, 18, RGBColor(198, 217, 237))
card(s, 0.82, 4.15, 5.85, 2.15, "Integrantes", "Jorge Gabriel Palacios Sales\nPablo Daniel Barillas Moreno\nRoberto Emiliano Otoniel Camposeco Torres\nAbby Sofia Donis Agreda", GOLD, 14)
card(s, 6.95, 4.15, 5.25, 2.15, "Curso y contexto", "CC3084 · Data Science\nSección 10 · Universidad del Valle de Guatemala\nSemestre II · 2026", CYAN, 14)
text(s, "Fuente de datos: competencia MAP de Kaggle", 0.84, 7.02, 6.5, 0.22, 9, RGBColor(164, 187, 211))

# 2 · Situación
s = base_slide("Una respuesta correcta no garantiza comprensión", "Situación problemática", 2)
card(s, 0.62, 1.76, 3.75, 4.75, "El reto educativo", "Dos estudiantes pueden seleccionar la misma opción y justificarla con razonamientos distintos. La corrección binaria pierde la señal pedagógica contenida en el lenguaje.", GOLD, 18)
card(s, 4.62, 1.76, 3.75, 4.75, "El reto analítico", "Las explicaciones son breves, desbalanceadas, dependientes de la pregunta y combinan lenguaje natural con símbolos matemáticos.", CYAN, 18)
card(s, 8.62, 1.76, 3.75, 4.75, "La oportunidad", "Clasificar categoría y concepción errónea permite priorizar retroalimentación y descubrir patrones sin sustituir el juicio docente.", GREEN, 18)

# 3 · Pregunta científica
s = base_slide("Pregunta científica y unidad de análisis", "Planteamiento", 3)
rect(s, 0.8, 1.78, 11.7, 1.65, NAVY, True, NAVY)
text(s, "¿Qué patrones textuales y contextuales distinguen las categorías de respuesta y las concepciones erróneas en las explicaciones estudiantiles?", 1.2, 2.08, 10.9, 1.02, 22, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
card(s, 0.8, 3.83, 3.55, 2.18, "Unidad", "Una respuesta estudiantil a una pregunta de opción múltiple.", CYAN, 16)
card(s, 4.72, 3.83, 3.55, 2.18, "Contexto", "QuestionText + MC_Answer + StudentExplanation.", GOLD, 16)
card(s, 8.65, 3.83, 3.55, 2.18, "Resultado", "Category y, cuando aplica, Misconception.", GREEN, 16)

# 4 · Objetivos
s = base_slide("Objetivos medibles del análisis", "Objetivos", 4)
card(s, 0.72, 1.72, 12.0, 1.18, "Objetivo general", "Caracterizar la estructura, calidad, desbalance y señales lingüísticas del conjunto MAP para fundamentar un pipeline de clasificación reproducible.", GOLD, 17)
card(s, 0.72, 3.2, 3.72, 2.8, "01 · Describir", "Cuantificar observaciones, variables, tipos, faltantes y duplicados.", CYAN, 17)
card(s, 4.8, 3.2, 3.72, 2.8, "02 · Explorar", "Comparar distribuciones, relaciones por pregunta y marcadores léxicos.", GREEN, 17)
card(s, 8.88, 3.2, 3.72, 2.8, "03 · Decidir", "Traducir los hallazgos en validación, preprocesamiento y modelos candidatos.", GOLD, 17)

# 5 · Datos
s = base_slide("Perfil verificado de los datos", "Descripción de los datos", 5)
metric(s, 0.7, 1.72, 2.2, "36,696", "RESPUESTAS", CYAN)
metric(s, 3.08, 1.72, 2.2, "15", "PREGUNTAS", GOLD)
metric(s, 5.46, 1.72, 2.2, "49", "OPCIONES ÚNICAS", GREEN)
metric(s, 7.84, 1.72, 2.2, "35", "MISCONCEPTIONS", RED)
metric(s, 10.22, 1.72, 2.2, "65", "TARGETS COMPUESTOS", BLUE)
card(s, 0.7, 3.3, 5.75, 2.65, "Esquema de entrenamiento", "row_id · QuestionId · QuestionText · MC_Answer · StudentExplanation · Category · Misconception", CYAN, 16)
card(s, 6.75, 3.3, 5.67, 2.65, "Integridad", "0 duplicados de fila completa. Existen 736 duplicados al excluir row_id y 970 textos repetidos; se controlan por grupos durante la validación.", GOLD, 16)

# 6 · Categoría
s = base_slide("La categoría combina dos decisiones", "Variable objetivo", 6)
picture(s, "distribucion_category.png", 0.75, 1.72, 11.85, 3.92)
card(s, 0.9, 5.77, 3.55, 0.82, "Resultado", "40.34% True_Correct.", GREEN, 12)
card(s, 4.86, 5.77, 3.55, 0.82, "Ruta", "Tipo × Correctness.", CYAN, 12)
card(s, 8.82, 5.77, 3.55, 0.82, "Uso", "Modelado jerárquico.", GOLD, 12)

# 7 · Limpieza
s = base_slide("Los nulos son estructurales, no aleatorios", "Limpieza y preprocesamiento", 7)
metric(s, 0.8, 1.78, 3.05, "26,836", "MISSING MISCONCEPTION", RED)
metric(s, 4.15, 1.78, 3.05, "73.13%", "DEL TOTAL", GOLD)
metric(s, 7.5, 1.78, 3.05, "100%", "COINCIDE CON NO-MISCONCEPTION", GREEN)
card(s, 0.8, 3.35, 5.55, 2.6, "Decisión", "Conservar el nulo como estado semántico. No imputar una concepción errónea donde la categoría indica que no existe.", CYAN, 17)
card(s, 6.7, 3.35, 5.55, 2.6, "Preprocesamiento", "Normalizar espacios y Unicode sin borrar números, operadores, signos o negaciones; esas marcas pueden cambiar el razonamiento.", GOLD, 17)

# 8 · Cola larga
s = base_slide("La etiqueta compuesta presenta una cola larga", "Desbalance", 8)
picture(s, "distribucion_misconceptions.png", 0.72, 1.7, 8.45, 4.58)
card(s, 9.45, 1.7, 2.85, 1.22, "32 clases", "tienen menos de 50 casos", RED, 14)
card(s, 9.45, 3.12, 2.85, 1.22, "24 clases", "tienen menos de 10 casos", GOLD, 14)
card(s, 9.45, 4.54, 2.85, 1.74, "Implicación", "Usar macro F1, MAP@3, pesos y análisis por clase.", CYAN, 14)

# 9 · Texto
s = base_slide("Explicaciones breves, pero heterogéneas", "Variables textuales", 9)
picture(s, "distribuciones_variables_texto.png", 0.72, 1.66, 8.42, 4.78)
metric(s, 9.44, 1.76, 2.82, "13", "MEDIANA DE PALABRAS", CYAN)
metric(s, 9.44, 3.18, 2.82, "15.36", "MEDIA DE PALABRAS", GOLD)
metric(s, 9.44, 4.60, 2.82, "177", "MÁXIMO", RED)

# 10 · Atípicos
s = base_slide("Los valores extremos contienen señal potencial", "Outliers", 10)
picture(s, "longitud_por_categoria.png", 0.72, 1.72, 8.52, 4.43)
card(s, 9.5, 1.72, 2.75, 2.02, "No eliminar automáticamente", "Una explicación extensa puede reflejar duda, justificación detallada o copia.", GOLD, 14)
card(s, 9.5, 4.02, 2.75, 2.13, "Sí auditar", "Revisar casos extremos, truncamiento y robustez del tokenizador.", CYAN, 14)

# 11 · Correlaciones
s = base_slide("Longitud y diversidad no significan lo mismo", "Relaciones", 11)
picture(s, "correlaciones_variables_texto.png", 0.72, 1.72, 8.58, 3.96)
card(s, 9.55, 1.72, 2.7, 1.55, "Relación clave", "Diversidad léxica vs. palabras: ≈ −0.59.", CYAN, 14)
card(s, 9.55, 3.52, 2.7, 2.16, "Lectura", "Las respuestas largas repiten más vocabulario; ninguna variable aislada explica la etiqueta.", GOLD, 14)

# 12 · Preguntas
s = base_slide("El contexto de la pregunta cambia el problema", "Cruces categóricos", 12)
picture(s, "cruces_por_pregunta.png", 0.72, 1.74, 8.82, 3.72)
card(s, 9.78, 1.74, 2.48, 1.38, "Rango", "35.7%–79.6% de True_Correct.", GOLD, 14)
card(s, 9.78, 3.37, 2.48, 2.09, "Cuidado", "QuestionId ayuda en preguntas conocidas, pero el ID crudo no generaliza a nuevas.", CYAN, 14)
text(s, "32 de 35 concepciones erróneas aparecen en una sola pregunta.", 0.95, 5.78, 11.0, 0.45, 18, NAVY, True, align=PP_ALIGN.CENTER)

# 13 · Marcadores
s = base_slide("Los marcadores léxicos son señales, no explicaciones", "Lenguaje", 13)
picture(s, "marcadores_lexicos.png", 0.72, 1.72, 8.45, 4.3)
card(s, 9.44, 1.72, 2.82, 1.5, "Ejemplos", "divide · multiply · denominator · because", CYAN, 14)
card(s, 9.44, 3.49, 2.82, 2.53, "Interpretación", "Su frecuencia cambia por categoría, pero debe evaluarse junto con pregunta, opción y explicación.", GOLD, 14)

# 14 · Investigación
s = base_slide("De líneas base interpretables a modelos contextuales", "Revisión bibliográfica", 14)
card(s, 0.72, 1.72, 3.7, 4.45, "1 · Baseline", "TF–IDF de palabras y caracteres con regresión logística y SVM.\n\nVentaja: costo bajo, transparencia y referencia fuerte para texto corto.", CYAN, 16)
card(s, 4.82, 1.72, 3.7, 4.45, "2 · Modelo principal", "BERT/DeBERTa con pregunta, opción y explicación; salida multitarea para categoría y misconception.\n\nVentaja: contexto semántico.", GOLD, 16)
card(s, 8.92, 1.72, 3.7, 4.45, "3 · Experimento", "MathBERT o adaptación de dominio, ensamble calibrado y ranking top–3.\n\nCondición: mejora consistente y costo justificable.", GREEN, 16)

# 15 · Hallazgos
s = base_slide("Cinco hallazgos que condicionan el modelado", "Síntesis", 15)
rich_text(s, [
    "Desbalance severo: 32 de 65 targets tienen menos de 50 ejemplos.",
    "Dependencia contextual: casi todas las misconceptions están ligadas a una pregunta.",
    "Nulos semánticos: 73.13% de Misconception representa ausencia legítima.",
    "Duplicación: 736 registros coinciden al excluir row_id.",
    "Texto corto y ruidoso: mediana 13 palabras; conservar notación matemática.",
], 0.92, 1.72, 11.35, 4.78, 20, INK, True)

# 16 · Validación
s = base_slide("Dos validaciones responden preguntas distintas", "Diseño experimental", 16)
card(s, 0.8, 1.72, 5.55, 3.25, "StratifiedKFold", "Escenario principal: rendimiento en nuevas respuestas a preguntas conocidas. Estratificar por etiqueta compuesta y agrupar duplicados.", CYAN, 17)
card(s, 6.7, 1.72, 5.55, 3.25, "GroupKFold por QuestionId", "Prueba de transferencia: extrapolación a preguntas no vistas. Reportarla por separado porque muchas etiquetas son específicas de pregunta.", GOLD, 17)
rect(s, 0.8, 5.25, 11.45, 0.88, NAVY, True, NAVY)
text(s, "Métricas: MAP@3 + macro F1 + desempeño por pregunta/clase + calibración", 1.05, 5.46, 10.95, 0.42, 17, WHITE, True, align=PP_ALIGN.CENTER)

# 17 · Repositorio
s = base_slide("Repositorio reproducible y entregables", "Trazabilidad", 17)
picture(s, "repo_qr.png", 0.9, 1.72, 2.52, 2.52)
text(s, "github.com/DanielBarillasM/\nProyecto-2_Grupo-1_DS_Sec-10", 0.78, 4.46, 3.05, 0.72, 13, BLUE, True, align=PP_ALIGN.CENTER)
card(s, 4.1, 1.72, 3.82, 4.35, "Evidencia", "• Notebook ejecutable\n• Datos versionados localmente\n• 7 figuras persistidas\n• Informe EDA en PDF/TeX\n• Revisión bibliográfica APA\n• README y documento del repo", CYAN, 16)
card(s, 8.32, 1.72, 3.82, 4.35, "Presentación", "• HTML navegable\n• PowerPoint editable\n• Rutas relativas a report/assets\n• Dependencias fijadas\n• Cifras verificadas contra train.csv", GOLD, 16)

# 18 · Cierre
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 0, 0.22, 7.5, CYAN)
picture(s, "logo_uvg.png", 11.25, 0.55, 0.92, 0.84)
text(s, "CONCLUSIÓN", 0.88, 0.82, 4.0, 0.35, 12, CYAN, True)
text(s, "Comprender el error requiere\nleer el razonamiento.", 0.88, 1.45, 10.8, 1.35, 33, WHITE, True)
text(s, "El EDA demuestra que la pregunta, la opción y la explicación deben modelarse juntas; el desbalance y la dependencia por pregunta exigen evaluación dual y resultados desagregados.", 0.9, 3.25, 10.6, 1.15, 20, RGBColor(205, 222, 239))
rect(s, 0.9, 5.05, 11.35, 0.95, RGBColor(12, 45, 80), True, RGBColor(12, 45, 80))
text(s, "Siguiente paso: baseline TF–IDF/SVM → Transformer contextual multitarea → ensamble calibrado", 1.2, 5.28, 10.75, 0.48, 17, GOLD, True, align=PP_ALIGN.CENTER)
text(s, "Gracias", 0.9, 6.82, 2.0, 0.35, 14, WHITE, True)

prs.core_properties.title = "MAP — Charting Student Math Misunderstandings"
prs.core_properties.subject = "Análisis exploratorio de datos"
prs.core_properties.author = "Grupo 1 — CC3084 Data Science, UVG"
prs.core_properties.keywords = "MAP, NLP, educación, matemáticas, EDA, UVG"
prs.save(OUTPUT)
print(f"Presentación generada: {OUTPUT} ({len(prs.slides)} diapositivas)")

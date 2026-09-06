"""Sincroniza el documento Word del repositorio con el estado verificado."""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "file" / "Presentacion_Repositorio_MAP.docx"
doc = Document(PATH)


def reemplazar_parrafo(indice, texto):
    doc.paragraphs[indice].text = texto


reemplazar_parrafo(1, "RESUMEN")
reemplazar_parrafo(
    3,
    "Este documento presenta cómo está organizado el repositorio, qué artefactos "
    "contiene y cómo reproducir el análisis. El notebook fue ejecutado desde cero "
    "con los CSV locales; sus 26 celdas de código finalizaron sin errores y las "
    "siete figuras de los entregables se regeneraron desde esas salidas.",
)
reemplazar_parrafo(46, "Estado de reproducibilidad")
reemplazar_parrafo(
    54,
    "Repositorio local y salidas verificadas de notebooks/01_analisis_exploratorio.ipynb.",
)

# Indicadores de portada interna.
t = doc.tables[1]
t.cell(0, 2).text = "35\nconcepciones erróneas"
t.cell(0, 3).text = "18\ndiapositivas PPTX/HTML"

# Árbol actualizado del repositorio.
doc.tables[3].cell(0, 0).text = """Proyecto-2_Grupo-1_DS_Sec-10/
├── data/
│   ├── train.csv
│   ├── test.csv
│   └── sample_submission.csv
├── instructions/
│   └── Proyecto 2. Análisis Exploratorio. 2026.pdf
├── notebooks/
│   ├── 01_analisis_exploratorio.ipynb
│   └── exportar_figuras.py
├── report/
│   ├── assets/
│   ├── informe_analisis_exploratorio_MAP.tex / .pdf
│   └── revision_bibliografica_algoritmos_MAP.tex / .pdf
├── presentation/
│   ├── presentacion_MAP.html
│   ├── presentacion_MAP.pptx / .pdf
│   └── generar_presentacion.py
├── file/
│   └── Presentacion_Repositorio_MAP.docx / .pdf
├── requirements.txt
├── LICENSE
└── README.md"""
arbol = doc.tables[3].cell(0, 0)
arbol.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
for parrafo in arbol.paragraphs:
    parrafo.paragraph_format.space_after = Pt(0)
    parrafo.paragraph_format.line_spacing = 0.92
    for run in parrafo.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

# Estado de componentes.
t = doc.tables[4]
t.cell(4, 2).text = "Informes, activos y presentaciones verificados localmente"
t.cell(5, 2).text = "Completo: contexto, resultados, reproducción y entregables"

# Flujo y entregables.
doc.tables[5].cell(4, 2).text = "PDF, HTML interactivo, PowerPoint editable y DOCX"
doc.tables[6].cell(0, 1).text = (
    "Regla de trazabilidad\nLas cifras se contrastaron con train.csv. El notebook "
    "se ejecutó en orden y las siete figuras se exportaron de sus salidas antes de "
    "regenerar los informes y la presentación."
)

t = doc.tables[7]
t.cell(1, 0).text = "01_analisis_exploratorio.ipynb"
t.cell(2, 0).text = "informes EDA y revisión bibliográfica (.pdf/.tex)"
t.cell(2, 2).text = "Rúbrica EDA, investigación APA y selección preliminar"
t.cell(3, 0).text = "presentacion_MAP.html / .pptx"
t.cell(3, 2).text = "Recorrido visual interactivo y PowerPoint editable"

# Hallazgos resumidos.
doc.tables[8].cell(0, 2).text = "13\npalabras de mediana"

# Reproducción.
t = doc.tables[10]
t.cell(4, 1).text = "Crear un entorno e instalar requirements.txt"
t.cell(4, 2).text = "Versiones fijadas para análisis y generación documental"

doc.tables[11].cell(0, 0).text = (
    "Ruta de datos\nEl notebook localiza data/ y data/raw desde la raíz o desde notebooks/."
)
doc.tables[11].cell(0, 1).text = (
    "Dependencias\nrequirements.txt fija el entorno de análisis y los generadores documentales."
)
doc.tables[12].cell(0, 1).text = (
    "Estado corregido\nRuta flexible, dependencias fijadas, notebook ejecutado sin errores, "
    "figuras sincronizadas y entregables compilados. Solo resta hacer commit y push "
    "cuando el equipo apruebe la versión final."
)

doc.core_properties.title = "Presentación del repositorio — MAP"
doc.core_properties.subject = "Trazabilidad y reproducibilidad del análisis exploratorio"
doc.core_properties.author = "Grupo 1 — CC3084 Data Science, UVG"
doc.core_properties.keywords = "MAP, NLP, EDA, repositorio, reproducibilidad, UVG"
doc.save(PATH)
print(f"Documento actualizado: {PATH}")
